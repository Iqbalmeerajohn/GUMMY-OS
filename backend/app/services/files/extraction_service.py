"""Text extraction from uploaded file bytes (M6; segments added for RAG 2.0).

Two entry points over the same parsers:

* :func:`extract_segments` returns the document as **structured segments** that
  carry where they came from — a PDF page, a Markdown heading, a span of CSV
  rows. This is what makes "Resume.pdf — page 2" possible: provenance is
  captured at extraction, because it cannot be recovered later.
* :func:`extract_text` flattens those segments back to a plain string. It is
  the original M6 signature and stays for callers that only want the text.

Before this, ``_extract_pdf`` joined pages with ``"

"`` and the page
numbers were gone by the time anything could use them.

Maps a MIME type to a decoder and returns plain text ready for chunking.
Parser libraries (``pypdf``, ``python-docx``, ``openpyxl``) are imported lazily
inside their branch — exactly like the optional embeddings providers — so they
are only required when a user actually uploads that format. Any extraction
failure is raised as :class:`ExtractionError` (an ``AppError``) so the service
layer can mark the file ``failed`` and capture the error.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass

from app.core.exceptions import AppError

logger = logging.getLogger(__name__)

# MIME types whose bytes are UTF-8 text and need no parser.
_PLAINTEXT_MIME_TYPES = frozenset({"text/plain", "text/markdown"})


# How many CSV rows travel together in one segment. Small enough that a chunk
# stays about one topic, large enough not to spend a whole embedding on a
# single row.
_CSV_ROWS_PER_SEGMENT = 25


@dataclass(frozen=True)
class DocumentSegment:
    """A run of text plus where in the document it came from.

    Every field beyond ``content`` is optional because provenance differs by
    format: PDFs have pages, Markdown has headings, CSVs have row ranges, and a
    plain text file has none of them. A segment with no provenance is still a
    valid segment — it just cites the filename alone.
    """

    content: str
    page: int | None = None
    section: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    # Repeated into every chunk cut from this segment. A CSV row means nothing
    # without its header, and chunk 7 of a spreadsheet would otherwise be a
    # grid of bare values.
    header: str | None = None


class ExtractionError(AppError):
    """Raised when a file's text could not be extracted."""

    def __init__(self, mime_type: str, detail: str) -> None:
        super().__init__(
            f"Could not extract text from {mime_type}: {detail}",
            code="file_extraction_failed",
            status_code=422,
        )


class UnsupportedFileTypeError(AppError):
    """Raised when a MIME type has no registered extractor."""

    def __init__(self, mime_type: str) -> None:
        super().__init__(
            f"Unsupported file type: {mime_type}",
            code="unsupported_file_type",
            status_code=415,
        )


def extract_segments(*, data: bytes, mime_type: str) -> list[DocumentSegment]:
    """Extract ``data`` as segments carrying their provenance.

    Same parsers and same failure modes as :func:`extract_text`; the difference
    is that page numbers, headings and row ranges survive.
    """
    if mime_type == "application/pdf":
        return _pdf_segments(data)
    if mime_type == "text/markdown":
        return _markdown_segments(_decode_text(data))
    if mime_type == "text/csv":
        return _csv_segments(data)
    if mime_type in _PLAINTEXT_MIME_TYPES:
        return _plain_segments(_decode_text(data))
    # DOCX/XLSX keep working, without structural attribution — outside the RAG
    # formats for this milestone, so they cite the filename only.
    text = extract_text(data=data, mime_type=mime_type)
    return _plain_segments(text)


def _plain_segments(text: str) -> list[DocumentSegment]:
    stripped = text.strip()
    return [DocumentSegment(content=stripped)] if stripped else []


def _pdf_segments(data: bytes) -> list[DocumentSegment]:
    """One segment per page, numbered from 1 as a reader would count them."""
    return [
        DocumentSegment(content=page_text.strip(), page=page_number)
        for page_number, page_text in enumerate(_pdf_pages(data), start=1)
        if page_text.strip()
    ]


# A Markdown ATX heading: one to six hashes, a space, then the title.
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def _markdown_segments(text: str) -> list[DocumentSegment]:
    """Split on headings, so each segment knows which section it belongs to.

    The heading itself is kept in the segment body — it is usually the single
    most retrievable line in the section, and dropping it would hide "Memory
    System" from a search for those words.
    """
    segments: list[DocumentSegment] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            segments.append(DocumentSegment(content=body, section=section))

    for line in text.splitlines():
        match = _HEADING_RE.match(line)
        if match:
            flush()
            buffer = [line]
            section = match.group(2).strip()
            continue
        buffer.append(line)
    flush()
    return segments


def _csv_segments(data: bytes) -> list[DocumentSegment]:
    """Row groups, each carrying the header and the row numbers it covers.

    Row numbers are the ones a spreadsheet shows: the header is row 1, so the
    first data row is row 2.
    """
    rows = list(csv.reader(io.StringIO(_decode_text(data))))
    if not rows:
        return []
    header_cells = rows[0]
    header = ", ".join(header_cells).strip()
    data_rows = rows[1:]
    if not data_rows:
        return [DocumentSegment(content=header, header=header, row_start=1, row_end=1)]

    segments: list[DocumentSegment] = []
    for start in range(0, len(data_rows), _CSV_ROWS_PER_SEGMENT):
        window = data_rows[start : start + _CSV_ROWS_PER_SEGMENT]
        body = "\n".join(", ".join(row) for row in window).strip()
        if not body:
            continue
        segments.append(
            DocumentSegment(
                content=body,
                header=header,
                row_start=start + 2,
                row_end=start + 1 + len(window),
            )
        )
    return segments


def _pdf_pages(data: bytes) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("application/pdf", "pypdf not installed") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ExtractionError("application/pdf", str(exc)) from exc


def extract_text(*, data: bytes, mime_type: str) -> str:
    """Extract plain text from ``data`` according to ``mime_type``.

    Returns possibly-empty text (an empty document is valid and simply yields
    zero chunks). Raises :class:`UnsupportedFileTypeError` for unknown types and
    :class:`ExtractionError` when a known parser fails.
    """
    if mime_type in _PLAINTEXT_MIME_TYPES:
        return _decode_text(data)
    if mime_type == "text/csv":
        return _extract_csv(data)
    if mime_type == "application/pdf":
        return _extract_pdf(data)
    if mime_type == (
        "application/vnd.openxmlformats-officedocument" ".wordprocessingml.document"
    ):
        return _extract_docx(data)
    if mime_type == (
        "application/vnd.openxmlformats-officedocument" ".spreadsheetml.sheet"
    ):
        return _extract_xlsx(data)
    raise UnsupportedFileTypeError(mime_type)


def _decode_text(data: bytes) -> str:
    """Decode bytes as UTF-8, replacing undecodable bytes (never raises)."""
    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    """Flatten CSV rows into newline-joined, comma-spaced text."""
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    return "\n".join(", ".join(row) for row in reader)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("application/pdf", "pypdf not installed") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ExtractionError("application/pdf", str(exc)) from exc
    return "\n\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document",
            "python-docx not installed",
        ) from exc
    try:
        document = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
    except Exception as exc:
        raise ExtractionError(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document",
            str(exc),
        ) from exc
    return "\n".join(paragraphs).strip()


def _extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError(
            "application/vnd.openxmlformats-officedocument" ".spreadsheetml.sheet",
            "openpyxl not installed",
        ) from exc
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
        lines: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(", ".join(cells))
        workbook.close()
    except Exception as exc:
        raise ExtractionError(
            "application/vnd.openxmlformats-officedocument" ".spreadsheetml.sheet",
            str(exc),
        ) from exc
    return "\n".join(lines).strip()
